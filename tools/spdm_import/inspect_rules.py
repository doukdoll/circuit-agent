"""Inspect merged/preprocessed SPDM workbooks locally; never modify the input."""

import argparse
from collections import Counter
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
import re
import sys
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.utils.exceptions import InvalidFileException


def json_value(value):
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def normalize(value, kind, epoch):
    if blank(value):
        return None
    if kind == "text":
        # Preserve internal whitespace, units, delimiters and conditions.
        return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    if kind == "sec_code":
        if not isinstance(value, str):
            raise ValueError("SEC 코드는 문자열이어야 함; 숫자 셀의 앞자리 0은 자동 복구하지 않음")
        value = value.strip()
        if not re.fullmatch(r"[0-9]{8}", value):
            raise ValueError("SEC 코드는 현재 입력 계약상 대시 없는 8자리 문자열이어야 함")
        return value
    if kind == "integer":
        if isinstance(value, bool):
            raise ValueError("정수 필드에 논리값이 들어 있음")
        if isinstance(value, int):
            return value
        if isinstance(value, float) and math.isfinite(value) and value.is_integer():
            return int(value)
        if isinstance(value, str) and re.fullmatch(r"[+-]?[0-9]+", value.strip()):
            return int(value.strip())
        raise ValueError("정수로 손실 없이 변환할 수 없음")
    if kind == "datetime":
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            value = from_excel(value, epoch)
        if not isinstance(value, datetime):
            raise ValueError("Excel 날짜/시각으로 변환할 수 없음")
        return value.isoformat()
    raise ValueError(f"지원하지 않는 타입: {kind}")


def read_mapping(path):
    config = json.loads(path.read_text(encoding="utf-8-sig"))
    fields = config.get("fields", [])
    names = [field["name"] for field in fields]
    if not fields or len(names) != len(set(names)):
        raise ValueError("매핑 필드가 없거나 내부 이름이 중복됨")
    for field in fields:
        if field["type"] not in {"text", "sec_code", "integer", "datetime"}:
            raise ValueError(f"지원하지 않는 타입: {field['type']}")
        if "header" not in field:
            raise ValueError("각 필드에 예상 header가 필요함")
        if "column" in field:
            column_index_from_string(field["column"])
    if any(name not in names for name in config.get("duplicate_key", [])):
        raise ValueError("중복 후보 키에 매핑되지 않은 필드가 있음")
    return config


def resolve_fields(headers, config):
    resolved = []
    for field in config["fields"]:
        if "column" in field:
            index = column_index_from_string(field["column"]) - 1
            if index >= len(headers) or headers[index] != field["header"]:
                raise ValueError(f"{field['column']}열의 예상 헤더 {field['header']!r}가 일치하지 않음")
        else:
            matches = [i for i, header in enumerate(headers) if header == field["header"]]
            if len(matches) != 1:
                raise ValueError(f"헤더 {field['header']!r}가 없거나 중복됨; column을 지정해야 함")
            index = matches[0]
        resolved.append((field, index))
    return resolved


def inspect_workbook(source, config, preview_limit=10):
    if preview_limit < 0:
        raise ValueError("preview는 0 이상이어야 함")
    # Read formulas literally so they are flagged rather than silently using stale caches.
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = workbook[config["sheet"]]
        rows = sheet.iter_rows()
        header_cells = next(rows, ())
        headers = [cell.value for cell in header_cells]
        fields = resolve_fields(headers, config)
        nonempty = Counter()
        selected_nonempty = Counter()
        unique_secs = set()
        keys = {}
        issues, duplicates, previews = [], [], []
        row_count = empty_rows = 0
        for row_number, cells in enumerate(rows, 2):
            values = [cell.value for cell in cells]
            if all(blank(value) for value in values):
                empty_rows += 1
                continue
            row_count += 1
            for index, value in enumerate(values):
                if not blank(value):
                    nonempty[index] += 1
            raw, cleaned = {}, {}
            for field, index in fields:
                value = values[index] if index < len(values) else None
                name = field["name"]
                raw[name] = json_value(value)
                try:
                    if index < len(cells) and cells[index].data_type in {"f", "e"}:
                        raise ValueError("수식 또는 Excel 오류 셀; 확정된 값으로 제공해야 함")
                    cleaned[name] = normalize(value, field["type"], workbook.epoch)
                    if cleaned[name] is None and field.get("required", False):
                        raise ValueError("필수 식별값이 비어 있음")
                    if cleaned[name] is not None:
                        selected_nonempty[name] += 1
                        if field["type"] == "sec_code":
                            unique_secs.add(cleaned[name])
                    if cleaned[name] is not None and "allowed" in field and cleaned[name] not in field["allowed"]:
                        issues.append({"row": row_number, "field": name, "severity": "warning",
                                       "reason": "예상 범위 밖의 값; 원문을 유지함"})
                except (ValueError, OverflowError, TypeError) as error:
                    cleaned[name] = None
                    issues.append({"row": row_number, "field": name, "severity": "error", "reason": str(error)})
            key_fields = config.get("duplicate_key", [])
            if key_fields and all(cleaned[name] is not None for name in key_fields):
                key = tuple(cleaned[name] for name in key_fields)
                if key in keys:
                    duplicates.append({"row": row_number, "first_row": keys[key]})
                else:
                    keys[key] = row_number
            if len(previews) < preview_limit:
                previews.append({"source_row": row_number, "raw": raw, "selected": cleaned})
        return {
            "source_file": source.name,
            "sheet": config["sheet"],
            "mapping_sha256": hashlib.sha256(json.dumps(config, ensure_ascii=False, sort_keys=True).encode()).hexdigest(),
            "data_rows": row_count,
            "empty_rows": empty_rows,
            "valid_unique_sec_codes": len(unique_secs),
            "columns": [{"column": get_column_letter(i + 1), "header": json_value(header),
                         "nonempty": nonempty[i], "empty": row_count - nonempty[i]} for i, header in enumerate(headers)],
            "selected_fields": [{"name": field["name"], "column": get_column_letter(index + 1),
                                 "header": field["header"], "type": field["type"],
                                 "nonempty": selected_nonempty[field["name"]]} for field, index in fields],
            "errors": sum(issue["severity"] == "error" for issue in issues),
            "warnings": sum(issue["severity"] == "warning" for issue in issues),
            "issues": issues,
            "duplicate_candidates": duplicates,
            "preview": previews,
        }
    finally:
        workbook.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description="SPDM 병합/정제 xlsx의 컬럼 선택과 최소 정제 점검 (DB 저장 없음)")
    parser.add_argument("input", type=Path)
    parser.add_argument("--mapping", type=Path, default=Path(__file__).with_name("mapping.json"))
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--preview", type=int, default=10)
    args = parser.parse_args(argv)
    try:
        if args.report.resolve() in {args.input.resolve(), args.mapping.resolve()}:
            raise ValueError("보고서 경로는 입력/매핑 파일과 달라야 함")
        config = read_mapping(args.mapping)
        report = inspect_workbook(args.input, config, args.preview)
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
        print(f"행={report['data_rows']}, 유효 SEC={report['valid_unique_sec_codes']}, "
              f"오류={report['errors']}, 경고={report['warnings']}, 중복후보={len(report['duplicate_candidates'])}")
        print(f"보고서: {args.report}")
        return 2 if report["errors"] else 0
    except (OSError, ValueError, KeyError, TypeError, BadZipFile, InvalidFileException) as error:
        print(f"입력/설정 오류: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
